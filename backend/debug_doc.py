import docx, os

# Check the ORIGINAL template
src = r'C:\Users\26370\Desktop\软件工程课程设计文档模板\5详细设计说明书.docx'
doc = docx.Document(src)

print('=== Original Template P10 ===')
p10 = doc.paragraphs[10]
print(f'repr: {repr(p10.text[:100])}')
print(f'hex: {p10.text[:50].encode("utf-8").hex()}')
print(f'runs: {len(p10.runs)}')
for i, r in enumerate(p10.runs):
    print(f'  run[{i}]: repr={repr(r.text[:50])}')

print()
print('=== Original Template P13 ===')
p13 = doc.paragraphs[13]
print(f'repr: {repr(p13.text[:200])}')
print(f'runs: {len(p13.runs)}')
for i, r in enumerate(p13.runs):
    print(f'  run[{i}]: repr={repr(r.text[:80])}')

print()
print('=== Original Template Table0 [2,1] ===')
t0 = doc.tables[0]
cell = t0.rows[2].cells[1]
print(f'repr: {repr(cell.text[:200])}')
print(f'runs in first para: {len(cell.paragraphs[0].runs)}')
for i, r in enumerate(cell.paragraphs[0].runs):
    print(f'  run[{i}]: repr={repr(r.text[:80])}')
