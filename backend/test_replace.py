import docx, os

src = r'C:\Users\26370\Desktop\软件工程课程设计文档模板\5详细设计说明书.docx'
doc = docx.Document(src)

# Test the replacement on P10
p10 = doc.paragraphs[10]
print(f'P10 para.text: {repr(p10.text)}')
print(f'P10 para.text hex: {p10.text.encode("utf-8").hex()}')

old = '2020年6月22日'
new = '2026年6月17日'

print(f'old in para.text: {old in p10.text}')
print(f'para.text.replace: {repr(p10.text.replace(old, new))}')

# Now try the actual replacement
if old in p10.text:
    new_text = p10.text.replace(old, new)
    print(f'New text: {repr(new_text)}')
    p10.clear()
    p10.add_run(new_text)
    print(f'After clear+add_run, para.text: {repr(p10.text)}')

# Check for non-breaking spaces
print()
print('Checking for special chars in P10:')
for i, ch in enumerate(p10.text):
    if ord(ch) > 127 or ch == '\xa0' or ch == '　':
        print(f'  char[{i}]: U+{ord(ch):04X} = {repr(ch)}')
