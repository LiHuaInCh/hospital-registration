import docx, os, sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

out = r'C:\Users\26370\Desktop\医院挂号导诊平台-课设文档'
doc = docx.Document(os.path.join(out, '7测试分析报告.docx'))

print('=== Doc 7: All paragraphs with text ===')
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t and len(t) > 3:
        print(f'P{i}: {t[:200]}')
print()

print('=== Doc 7: Tables ===')
for ti, t in enumerate(doc.tables):
    print(f'--- Table {ti} ({len(t.rows)}r x {len(t.columns)}c) ---')
    for ri, row in enumerate(t.rows):
        cells = [c.text[:80] for c in row.cells]
        print(f'  Row{ri}: {" | ".join(cells)}')
    print()
