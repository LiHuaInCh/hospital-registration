import docx, os, sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

out = r'C:\Users\26370\Desktop\医院挂号导诊平台-课设文档'
doc = docx.Document(os.path.join(out, '8个人总结.docx'))

# Check all placeholders are gone
full_text = ' '.join([p.text for p in doc.paragraphs])
for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            full_text += ' ' + c.text

remaining = []
for p in ['****', '**系统', '2020年', '软件171', '***']:
    if p in full_text:
        remaining.append(p)
print(f'Placeholders remaining: {remaining if remaining else "NONE"}')

# Check Table0 (document list)
print('\n=== Table 0: Document List ===')
t0 = doc.tables[0]
for ri, row in enumerate(t0.rows):
    cells = [c.text[:80] for c in row.cells]
    print(f'  Row{ri}: {" | ".join(cells)}')

# Check Table1 (personal summaries)
print(f'\n=== Table 1: Personal Summaries ({len(doc.tables[1].rows)} rows) ===')
t1 = doc.tables[1]
for ri, row in enumerate(t1.rows):
    txt = row.cells[0].text[:120]
    print(f'  Row{ri}: {txt}...' if len(row.cells[0].text) > 120 else f'  Row{ri}: {txt}')
