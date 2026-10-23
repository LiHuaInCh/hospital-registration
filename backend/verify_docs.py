import docx, os, sys

# Force UTF-8 output
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

out = r'C:\Users\26370\Desktop\医院挂号导诊平台-课设文档'
for fname in ['5详细设计说明书.docx','6软件测试计划.docx','7测试分析报告.docx','8个人总结.docx']:
    fp = os.path.join(out, fname)
    doc = docx.Document(fp)
    full_text = ' '.join([p.text for p in doc.paragraphs])
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                full_text += ' ' + c.text

    remaining = []
    for p in ['****', '**系统', '2020年', '软件171', '***']:
        if p in full_text:
            count = full_text.count(p)
            remaining.append(f'{p}({count})')

    status = 'CLEAN' if not remaining else f'HAS: {remaining}'
    print(f'{fname}: {status}')
    # Print the title paragraph
    for para in doc.paragraphs:
        t = para.text.strip()
        if t and len(t) > 5:
            print(f'  First content: {t[:120]}')
            break
    print()
