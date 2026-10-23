import docx, os, shutil, sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# Import the actual functions from fill_docs
sys.path.insert(0, r'C:\Users\26370\Desktop\hospital-registration\backend')
from fill_docs import (
    replace_all_text, replace_ordered, COMMON_REPLACEMENTS,
    fill_table_members, fill_doc5_content, post_process_remaining_stars
)

TEMPLATE_DIR = r'C:\Users\26370\Desktop\软件工程课程设计文档模板'
src = os.path.join(TEMPLATE_DIR, '5详细设计说明书.docx')
dst = r'C:\Users\26370\Desktop\hospital-registration\backend\test_full_output.docx'
shutil.copy(src, dst)

doc = docx.Document(dst)

print(f'BEFORE: P10 = {repr(doc.paragraphs[10].text)}')

replace_ordered(doc, COMMON_REPLACEMENTS)
print(f'After replace_ordered: P10 = {repr(doc.paragraphs[10].text)}')

fill_table_members(doc)
print(f'After fill_table_members: P10 = {repr(doc.paragraphs[10].text)}')

fill_doc5_content(doc)
print(f'After fill_doc5_content: P10 = {repr(doc.paragraphs[10].text)}')

post_process_remaining_stars(doc)
print(f'After post_process: P10 = {repr(doc.paragraphs[10].text)}')

doc.save(dst)
print(f'FILE SAVED')

# Re-read and check
doc2 = docx.Document(dst)
print(f'RE-READ: P10 = {repr(doc2.paragraphs[10].text)}')
print(f'RE-READ: P13 = {repr(doc2.paragraphs[13].text[:150])}')
