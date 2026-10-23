import docx, os, sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

out = r'C:\Users\26370\Desktop\医院挂号导诊平台-课设文档'
doc = docx.Document(os.path.join(out, '5详细设计说明书.docx'))

print('=== Checking for 2020年 ===')
for i, p in enumerate(doc.paragraphs):
    if '2020' in p.text:
        print(f'P{i}: ...{p.text[max(0,p.text.find("2020")-20):p.text.find("2020")+30]}...')

for ti, t in enumerate(doc.tables):
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            if '2020' in cell.text:
                idx = cell.text.find('2020')
                print(f'T{ti}[{ri},{ci}]: ...{cell.text[max(0,idx-20):idx+30]}...')

print()
print('=== Checking for 软件171 ===')
for i, p in enumerate(doc.paragraphs):
    if '软件171' in p.text:
        print(f'P{i}: ...{p.text[max(0,p.text.find("软件171")-20):p.text.find("软件171")+30]}...')

for ti, t in enumerate(doc.tables):
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            if '软件171' in cell.text:
                idx = cell.text.find('软件171')
                print(f'T{ti}[{ri},{ci}]: ...{cell.text[max(0,idx-20):idx+30]}...')
