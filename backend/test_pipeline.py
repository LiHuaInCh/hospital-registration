import docx, os, shutil

TEMPLATE_DIR = r'C:\Users\26370\Desktop\软件工程课程设计文档模板'
src = os.path.join(TEMPLATE_DIR, '5详细设计说明书.docx')
dst = r'C:\Users\26370\Desktop\hospital-registration\backend\test_output.docx'
shutil.copy(src, dst)

doc = docx.Document(dst)

SYS_NAME = '医院门诊挂号与导诊服务平台'
LEADER = '成员A'
DATE_FULL = '2026年6月17日'
CLASS_NAME = '软件201班'
CLASS_SHORT = '软件201'
MEMBERS_STR = '成员A、成员B、成员C、成员D、成员E、成员F、成员G'

def replace_all_text(doc, old, new):
    for para in doc.paragraphs:
        if old in para.text:
            first_run = para.runs[0] if para.runs else None
            font_info = None
            if first_run:
                font_info = {
                    'name': first_run.font.name,
                    'size': first_run.font.size,
                    'bold': first_run.font.bold,
                }
            new_text = para.text.replace(old, new)
            para.clear()
            run = para.add_run(new_text)
            if font_info:
                try:
                    if font_info['name']: run.font.name = font_info['name']
                    if font_info['size']: run.font.size = font_info['size']
                    if font_info['bold'] is not None: run.font.bold = font_info['bold']
                except Exception: pass
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if old in para.text:
                        first_run = para.runs[0] if para.runs else None
                        font_info = None
                        if first_run:
                            font_info = {
                                'name': first_run.font.name,
                                'size': first_run.font.size,
                                'bold': first_run.font.bold,
                            }
                        new_text = para.text.replace(old, new)
                        para.clear()
                        run = para.add_run(new_text)
                        if font_info:
                            try:
                                if font_info['name']: run.font.name = font_info['name']
                                if font_info['size']: run.font.size = font_info['size']
                                if font_info['bold'] is not None: run.font.bold = font_info['bold']
                            except Exception: pass

COMMON = [
    ('****系统', SYS_NAME),
    ('****', SYS_NAME),
    ('**系统', SYS_NAME),
    ('***', LEADER),
    ('2020年6月22日', DATE_FULL),
    ('2020年 6月 22日', DATE_FULL),
    ('软件171班', CLASS_NAME),
    ('软件171', CLASS_SHORT),
    ('2017******', '2021******'),
    ('**、**、**、**、**、**', MEMBERS_STR),
]

# Check P10 before any changes
print(f'BEFORE: P10 = {repr(doc.paragraphs[10].text)}')
print(f'BEFORE: P13 = {repr(doc.paragraphs[13].text[:100])}')

# Apply replacements one by one
for i, (old, new) in enumerate(COMMON):
    replace_all_text(doc, old, new)
    if i == 4 or i == 5:  # date replacements
        print(f'After step {i} ({old[:20]}...): P10 = {repr(doc.paragraphs[10].text)}')

print()
print(f'AFTER ALL: P10 = {repr(doc.paragraphs[10].text)}')
print(f'AFTER ALL: P13 = {repr(doc.paragraphs[13].text[:150])}')

# Check Table0
t0 = doc.tables[0]
print(f'AFTER ALL: T0[2,1] = {repr(t0.rows[2].cells[1].text[:100])}')
