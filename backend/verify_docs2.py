import docx, os, sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

out = r'C:\Users\26370\Desktop\医院挂号导诊平台-课设文档'

print('=' * 60)
print('DOCUMENT 5: 详细设计说明书')
print('=' * 60)
doc = docx.Document(os.path.join(out, '5详细设计说明书.docx'))

# Check key paragraphs
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t and len(t) > 5:
        print(f'P{i}: {t[:150]}')
print()

# Check tables
for ti, t in enumerate(doc.tables):
    print(f'--- Table {ti} ({len(t.rows)}r x {len(t.columns)}c) ---')
    for ri, row in enumerate(t.rows):
        cells = [c.text[:60] for c in row.cells]
        print(f'  Row{ri}: {" | ".join(cells)}')
    print()

print()
print('=' * 60)
print('DOCUMENT 6: 软件测试计划')
print('=' * 60)
doc = docx.Document(os.path.join(out, '6软件测试计划.docx'))
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t and len(t) > 5:
        print(f'P{i}: {t[:150]}')
        if i > 50:
            print('... (truncated)')
            break
print()

for ti, t in enumerate(doc.tables):
    print(f'--- Table {ti} ({len(t.rows)}r x {len(t.columns)}c) ---')
    for ri, row in enumerate(t.rows):
        cells = [c.text[:60] for c in row.cells]
        print(f'  Row{ri}: {" | ".join(cells)}')
    print()

print()
print('=' * 60)
print('DOCUMENT 8: 个人总结')
print('=' * 60)
doc = docx.Document(os.path.join(out, '8个人总结.docx'))
for ti, t in enumerate(doc.tables):
    print(f'--- Table {ti} ({len(t.rows)}r x {len(t.columns)}c) ---')
    for ri, row in enumerate(t.rows):
        cells = [c.text[:80] for c in row.cells]
        print(f'  Row{ri}: {" | ".join(cells)}')
    print()
